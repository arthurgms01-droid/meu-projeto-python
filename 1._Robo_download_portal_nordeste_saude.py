from __future__ import annotations

import base64
import csv
import logging
import os
import queue
import re
import sys
import threading
import time
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from tkinter import BOTH, END, LEFT, RIGHT, X, Button, Entry, Frame, Label, StringVar, Tk, filedialog, messagebox
from tkinter.scrolledtext import ScrolledText
from tkinter.ttk import Combobox
from urllib.parse import urljoin

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import (
    Browser,
    BrowserContext,
    Download,
    Locator,
    Page,
    Response,
    TimeoutError as PlaywrightTimeoutError,
    sync_playwright,
)


# =============================================================================
# CONFIGURAÇÃO PRINCIPAL
# =============================================================================

URL_SITE = "https://nordestesaude.com.br/"
URL_PREFEITURA = "http://sefazweb.camacari.ba.gov.br/nfse/consultarAutenticidadeNFSe.tela"

NOME_BASE_PLANILHA = "NORDEST_CONTRATOS_DOWNLOAD"
NOME_PASTA_DOWNLOADS = "Downloads_Nordeste_Saude"
NOME_PASTA_LOGS = "Logs_Nordeste_Saude"

XPATH_SOU_EMPRESA = "/html/body/header/div[1]/div/div[2]/div/nav[1]/ul/li[3]/a"
XPATH_TIPO_ACESSO = "/html/body/main/div[2]/div/form/div/div[2]/div[1]/select"
XPATH_USUARIO = "/html/body/main/div[2]/div/form/div/div[2]/div[2]/input"
XPATH_SENHA = "/html/body/main/div[2]/div/form/div/div[2]/div[3]/input"
XPATH_ENTRAR = "/html/body/main/div[2]/div/form/div/div[3]/button"
XPATH_MEUS_SERVICOS = "/html/body/header/header/div[1]/div/div/div[2]/div/div[1]/div[1]"
XPATH_TABELA_BOLETOS = "/html/body/main/div[3]/div[3]/div[1]/table"
XPATH_VERIFICAR_PREFEITURA = "/html/body/form/div[2]/center/table[2]/tbody/tr/td[3]/table/tbody/tr/td[2]/a"
XPATH_BOLETAS_MODAL_CSV = "/html/body/main/div[3]/div[3]/div[2]/div/div/div[2]/div/div/div/div[1]/button/i"
XPATH_BOLETAS_MODAL_PDF = "/html/body/main/div[3]/div[3]/div[2]/div/div/div[2]/div/div/div/div[2]/button/i"

TIMEOUT_PADRAO_MS = 30_000
TIMEOUT_CARREGAMENTO_MS = 60_000
TIMEOUT_DOWNLOAD_MS = 90_000
TIMEOUT_LOGIN_MANUAL_SEGUNDOS = 180
MAXIMO_TENTATIVAS_CONTA = 2
MAXIMO_TENTATIVAS_DOWNLOAD = 2
EXECUTAR_VISIVEL = True
PREFERIR_MICROSOFT_EDGE = True

SCRIPT_CAPTURA_BLOB = r"""
(() => {
    if (window.__roboCapturaBlobInstalada) return;
    window.__roboCapturaBlobInstalada = true;
    window.__roboBlobsCriados = new Map();
    const criarURLOriginal = URL.createObjectURL.bind(URL);
    URL.createObjectURL = function(objeto) {
        const url = criarURLOriginal(objeto);
        try {
            if (objeto instanceof Blob) window.__roboBlobsCriados.set(url, objeto);
        } catch (_) {}
        return url;
    };
})();
"""


@dataclass(frozen=True)
class Credencial:
    linha_excel: int
    usuario: str
    senha: str
    inscricao_municipal: str
    contrato: str = ""
    codigo: str = ""
    vigencia: str = ""


@dataclass(frozen=True)
class LinhaBoletoPortal:
    codigo: str
    numero: str
    vencimento: str
    valor: str


@dataclass
class ResultadoConta:
    linha_excel: int
    usuario_mascarado: str
    contrato: str
    status: str = "PENDENTE"
    codigos_localizados: int = 0
    notas_baixadas: int = 0
    boletos_baixados: int = 0
    relatorios_baixados: int = 0
    ignorados: int = 0
    indisponiveis: int = 0
    falhas: int = 0
    detalhe: str = ""


@dataclass
class DadosExecucao:
    registros_portal: list[dict[str, str]] = field(default_factory=list)
    registros_notas: list[dict[str, str]] = field(default_factory=list)


class FilaLogHandler(logging.Handler):
    def __init__(self, fila: queue.Queue[str]) -> None:
        super().__init__()
        self.fila = fila

    def emit(self, record: logging.LogRecord) -> None:
        try:
            self.fila.put_nowait(self.format(record))
        except Exception:
            pass


class FiltroLogErros(logging.Filter):
    def filter(self, record: logging.LogRecord) -> bool:
        mensagem = record.getMessage()
        return (
            record.levelno >= logging.WARNING
            or "DIAGNOSTICO_PDF" in mensagem
            or "DIAGNOSTICO_BLOB" in mensagem
            or "ETAPA_ROBO" in mensagem
        )


def formatar_contexto_etapa(contexto: dict[str, object]) -> str:
    partes = [f"{chave.upper()}={texto(valor)}" for chave, valor in contexto.items() if texto(valor)]
    return " | " + " | ".join(partes) if partes else ""


@contextmanager
def etapa_registrada(logger: logging.Logger, etapa: str, **contexto: object):
    """Registra início, sucesso, duração e traceback de uma etapa do robô."""
    detalhes = formatar_contexto_etapa(contexto)
    inicio = time.monotonic()
    logger.info("ETAPA_ROBO | STATUS=INICIO | ETAPA=%s%s", etapa, detalhes)
    try:
        yield
    except Exception as exc:
        duracao = time.monotonic() - inicio
        logger.exception(
            "ETAPA_ROBO | STATUS=ERRO | ETAPA=%s | DURACAO_SEG=%.2f | TIPO_ERRO=%s | ERRO=%s%s",
            etapa,
            duracao,
            type(exc).__name__,
            exc,
            detalhes,
        )
        raise
    else:
        duracao = time.monotonic() - inicio
        logger.info(
            "ETAPA_ROBO | STATUS=SUCESSO | ETAPA=%s | DURACAO_SEG=%.2f%s",
            etapa,
            duracao,
            detalhes,
        )


def pasta_do_script() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def normalizar(valor: object) -> str:
    valor_str = unicodedata.normalize("NFKD", texto(valor))
    valor_str = "".join(ch for ch in valor_str if not unicodedata.combining(ch))
    return re.sub(r"[^a-zA-Z0-9]+", "_", valor_str.lower()).strip("_")


def nome_seguro(valor: str, padrao: str = "arquivo") -> str:
    valor = unicodedata.normalize("NFKD", valor)
    valor = "".join(ch for ch in valor if not unicodedata.combining(ch))
    valor = re.sub(r"[^a-zA-Z0-9._-]+", "_", valor).strip("._-")
    return valor[:150] or padrao


def data_para_nome(valor: str) -> str:
    encontrado = re.search(r"(\d{2})/(\d{2})/(\d{4})", texto(valor))
    if not encontrado:
        return "NAO_INFORMADO"
    return "-".join(encontrado.groups())


def valor_para_nome(valor: str) -> str:
    encontrado = re.search(r"(\d{1,3}(?:\.\d{3})*,\d{2}|\d+,\d{2})", texto(valor))
    return encontrado.group(1) if encontrado else "NAO_INFORMADO"


def mascarar_usuario(usuario: str) -> str:
    if "@" in usuario:
        inicio, dominio = usuario.split("@", 1)
        return f"{inicio[:2]}***@{dominio}"
    if len(usuario) <= 4:
        return "***"
    return f"{usuario[:2]}***{usuario[-2:]}"


def apenas_digitos(valor: str) -> str:
    return re.sub(r"\D", "", valor)


def padronizar_vigencia(valor: str) -> str:
    digitos = apenas_digitos(valor)
    if not digitos:
        return ""
    return f"VIG {int(digitos):02d}"


def procurar_planilha_existente(base: Path) -> Path | None:
    candidatos: list[Path] = []
    for extensao in (".xlsx", ".xlsm"):
        candidatos.extend(base.glob(f"{NOME_BASE_PLANILHA}*{extensao}"))
    candidatos = [p for p in candidatos if not p.name.startswith("~$")]
    if not candidatos:
        return None
    controles_finais = [p for p in candidatos if "MODELO" not in p.stem.upper()]
    grupo_escolhido = controles_finais or candidatos
    return max(grupo_escolhido, key=lambda p: p.stat().st_mtime)


def localizar_planilha(base: Path) -> Path:
    planilha = procurar_planilha_existente(base)
    if planilha is not None:
        return planilha
    modelo = criar_planilha_modelo(base)
    raise FileNotFoundError(f"Planilha não encontrada. Foi criado o modelo: {modelo}")


def criar_planilha_modelo(base: Path) -> Path:
    caminho = base / f"{NOME_BASE_PLANILHA}_MODELO.xlsx"
    if caminho.exists():
        return caminho
    wb = Workbook()
    ws = wb.active
    ws.title = "CONTROLE"
    ws.append(["USUARIO", "SENHA", "INSCRICAO_MUNICIPAL_PRESTADOR", "CONTRATO", "CODIGO", "VIGENCIA", "ATIVO"])
    ws.append(["usuario_exemplo", "senha_exemplo", "123456", "contrato_exemplo", "95774", "20", "SIM"])
    ws.freeze_panes = "A2"
    for coluna, largura in zip("ABCDEFG", (25, 22, 34, 22, 14, 14, 12)):
        ws.column_dimensions[coluna].width = largura
    wb.save(caminho)
    return caminho


def escolher_coluna(cabecalhos: dict[str, int], aliases: tuple[str, ...], obrigatoria: bool) -> int | None:
    for alias in aliases:
        if alias in cabecalhos:
            return cabecalhos[alias]
    if obrigatoria:
        raise ValueError("Coluna obrigatória ausente. Nomes aceitos: " + ", ".join(aliases))
    return None


def ler_credenciais(planilha: Path, vigencia_escolhida: str) -> list[Credencial]:
    wb = load_workbook(planilha, read_only=True, data_only=True)
    try:
        ws = next((aba for aba in wb.worksheets if aba.max_row >= 2), wb.active)
        cabecalhos = {
            normalizar(celula.value): indice
            for indice, celula in enumerate(ws[1], start=1)
            if texto(celula.value)
        }
        col_usuario = escolher_coluna(cabecalhos, ("usuario", "login", "email", "cpf", "cnpj", "usuario_portal"), True)
        col_senha = escolher_coluna(cabecalhos, ("senha", "password", "senha_portal"), True)
        col_inscricao = escolher_coluna(
            cabecalhos,
            ("inscricao_municipal_prestador", "inscricao_municipal", "inscricao", "im_prestador"),
            True,
        )
        col_contrato = escolher_coluna(cabecalhos, ("contrato", "numero_contrato", "codigo_contrato"), False)
        col_codigo = escolher_coluna(cabecalhos, ("codigo", "codigo_numero", "codigo_empresa"), False)
        col_vigencia = escolher_coluna(cabecalhos, ("vigencia", "vig", "dia_vencimento"), False)
        col_ativo = escolher_coluna(cabecalhos, ("ativo", "processar", "status"), False)

        credenciais: list[Credencial] = []
        for linha in range(2, ws.max_row + 1):
            usuario = texto(ws.cell(linha, col_usuario).value)
            senha = texto(ws.cell(linha, col_senha).value)
            inscricao = texto(ws.cell(linha, col_inscricao).value)
            ativo = texto(ws.cell(linha, col_ativo).value) if col_ativo else "SIM"
            codigo = texto(ws.cell(linha, col_codigo).value) if col_codigo else ""
            vigencia = padronizar_vigencia(texto(ws.cell(linha, col_vigencia).value)) if col_vigencia else ""

            if normalizar(ativo) in {"nao", "n", "0", "inativo", "ignorar"}:
                continue
            if not usuario and not senha and not inscricao:
                continue
            if vigencia_escolhida != "TODAS" and vigencia and vigencia != vigencia_escolhida:
                continue
            if not usuario or not senha or not inscricao:
                raise ValueError(
                    f"Linha {linha}: USUARIO, SENHA e INSCRICAO_MUNICIPAL_PRESTADOR são obrigatórios."
                )
            credenciais.append(
                Credencial(
                    linha_excel=linha,
                    usuario=usuario,
                    senha=senha,
                    inscricao_municipal=inscricao,
                    contrato=texto(ws.cell(linha, col_contrato).value) if col_contrato else "",
                    codigo=codigo,
                    vigencia=vigencia,
                )
            )
        if not credenciais:
            raise ValueError(f"Nenhuma linha ativa foi encontrada para {vigencia_escolhida}.")
        return credenciais
    finally:
        wb.close()


def configurar_logger(pasta_logs: Path, fila: queue.Queue[str]) -> tuple[logging.Logger, Path, Path]:
    pasta_logs.mkdir(parents=True, exist_ok=True)
    sufixo = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho = pasta_logs / f"execucao_{sufixo}.log"
    caminho_erros = pasta_logs / f"erros_{sufixo}.log"
    logger = logging.getLogger(f"nordeste_{sufixo}")
    logger.setLevel(logging.INFO)
    logger.propagate = False
    formato = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s", "%d/%m/%Y %H:%M:%S")
    arquivo = logging.FileHandler(caminho, encoding="utf-8")
    arquivo.setFormatter(formato)
    logger.addHandler(arquivo)

    arquivo_erros = logging.FileHandler(caminho_erros, encoding="utf-8")
    arquivo_erros.setFormatter(formato)
    arquivo_erros.addFilter(FiltroLogErros())
    logger.addHandler(arquivo_erros)
    interface = FilaLogHandler(fila)
    interface.setFormatter(formato)
    logger.addHandler(interface)
    return logger, caminho, caminho_erros


def aguardar_carregamento(page: Page, pausa_ms: int = 1_000) -> None:
    try:
        page.wait_for_load_state("domcontentloaded", timeout=TIMEOUT_CARREGAMENTO_MS)
    except PlaywrightTimeoutError:
        pass
    try:
        page.wait_for_load_state("networkidle", timeout=15_000)
    except PlaywrightTimeoutError:
        pass
    page.wait_for_timeout(pausa_ms)


def abrir_portal_empresa(contexto: BrowserContext, logger: logging.Logger) -> Page:
    page = contexto.new_page()
    page.goto(URL_SITE, wait_until="domcontentloaded", timeout=TIMEOUT_CARREGAMENTO_MS)
    paginas_antes = set(contexto.pages)
    link = page.locator(f"xpath={XPATH_SOU_EMPRESA}")
    link.wait_for(state="visible", timeout=TIMEOUT_PADRAO_MS)
    href = link.get_attribute("href")
    link.click()
    page.wait_for_timeout(1_500)
    novas = [p for p in contexto.pages if p not in paginas_antes]
    if novas:
        page = novas[-1]
    elif page.url == URL_SITE and href:
        page.goto(urljoin(URL_SITE, href), wait_until="domcontentloaded", timeout=TIMEOUT_CARREGAMENTO_MS)
    aguardar_carregamento(page)
    logger.info("Portal Empresa aberto: %s", page.url)
    return page


def selecionar_empresa(page: Page) -> None:
    select = page.locator(f"xpath={XPATH_TIPO_ACESSO}")
    select.wait_for(state="visible", timeout=TIMEOUT_PADRAO_MS)
    try:
        select.select_option(label="Empresa")
        return
    except Exception:
        pass
    opcoes = select.locator("option").all()
    for opcao in opcoes:
        if normalizar(opcao.inner_text()) == "empresa":
            select.select_option(value=opcao.get_attribute("value"))
            return
    raise RuntimeError("A opção 'Empresa' não foi encontrada na lista Tipo de acesso.")


def login_concluido(page: Page, url_anterior: str) -> bool:
    try:
        campo_senha_visivel = page.locator(f"xpath={XPATH_SENHA}").is_visible(timeout=500)
    except Exception:
        campo_senha_visivel = False
    return page.url != url_anterior and not campo_senha_visivel


def efetuar_login(page: Page, credencial: Credencial, logger: logging.Logger) -> None:
    selecionar_empresa(page)
    page.locator(f"xpath={XPATH_USUARIO}").fill(credencial.usuario)
    page.locator(f"xpath={XPATH_SENHA}").fill(credencial.senha)
    url_anterior = page.url
    page.locator(f"xpath={XPATH_ENTRAR}").click()
    logger.info("Linha %s: login enviado como Empresa.", credencial.linha_excel)

    limite = time.monotonic() + TIMEOUT_LOGIN_MANUAL_SEGUNDOS
    while time.monotonic() < limite:
        if login_concluido(page, url_anterior):
            aguardar_carregamento(page, 1_500)
            logger.info("Linha %s: login concluído.", credencial.linha_excel)
            return
        page.wait_for_timeout(1_000)
    raise TimeoutError("O portal não confirmou o login dentro do tempo limite.")


def localizar_opcao_segunda_via(page: Page) -> Locator | None:
    padrao = re.compile(r"(?:segunda|2\s*[ªaºo]?)\s*via\s*(?:de\s*)?boletos?", re.I)
    for tipo in ("link", "button"):
        try:
            loc = page.get_by_role(tipo, name=padrao).first
            if loc.count() and loc.is_visible(timeout=1_000):
                return loc
        except Exception:
            pass
    try:
        loc = page.get_by_text(padrao).first
        if loc.count() and loc.is_visible(timeout=1_000):
            return loc
    except Exception:
        pass
    return None


def tabela_boletos(page: Page) -> Locator:
    """Localiza a grade pelos cabeçalhos, independentemente das divs do layout."""
    cabecalho_codigo = page.locator("th", has_text=re.compile(r"c[oó]digos?", re.I))
    return page.locator("main table:has(thead):has(tbody)").filter(has=cabecalho_codigo).first


def aviso_sem_registros(page: Page) -> bool:
    aviso = page.get_by_text(
        re.compile(r"Nenhum (?:registro|boleto) localizado", re.I)
    ).first
    return aviso.count() > 0 and aviso.is_visible(timeout=500)


def navegar_segunda_via(page: Page, logger: logging.Logger) -> None:
    meus_servicos = page.locator(f"xpath={XPATH_MEUS_SERVICOS}")
    meus_servicos.wait_for(state="visible", timeout=TIMEOUT_PADRAO_MS)
    meus_servicos.click()
    page.wait_for_timeout(1_000)
    opcao = localizar_opcao_segunda_via(page)
    if opcao is None:
        raise RuntimeError(
            "O menu 'Meus Serviços' foi aberto, mas a opção '2ª Via de Boletos' não foi localizada."
        )
    opcao.click()
    aguardar_carregamento(page, 1_500)
    page.locator("main").get_by_text(re.compile(r"Segunda Via de Boletos", re.I)).first.wait_for(
        state="visible", timeout=TIMEOUT_CARREGAMENTO_MS
    )
    limite = time.monotonic() + 15
    while time.monotonic() < limite:
        tabela = tabela_boletos(page)
        if tabela.count() and tabela.is_visible(timeout=500):
            logger.info("Tela Segunda Via de Boletos carregada com tabela.")
            return
        if aviso_sem_registros(page):
            logger.info("Tela Segunda Via de Boletos carregada sem registros.")
            return
        page.wait_for_timeout(500)
    raise RuntimeError("Tela de boletos aberta, mas sem tabela nem aviso de ausência de registros.")


def cabecalhos_fisicos(tabela: Locator, quantidade_colunas: int) -> list[str]:
    cabecalhos: list[str] = []
    try:
        celulas = tabela.locator("thead tr").last.locator("th")
        for indice in range(celulas.count()):
            th = celulas.nth(indice)
            nome = texto(th.inner_text()) or f"COLUNA_{len(cabecalhos) + 1:02d}"
            colspan = int(th.get_attribute("colspan") or "1")
            for repeticao in range(colspan):
                cabecalhos.append(nome if repeticao == 0 else f"{nome}_{repeticao + 1}")
    except Exception:
        cabecalhos = []
    while len(cabecalhos) < quantidade_colunas:
        cabecalhos.append(f"COLUNA_{len(cabecalhos) + 1:02d}")
    return cabecalhos[:quantidade_colunas]


def nome_coluna_unico(nome: str, existentes: set[str]) -> str:
    base = nome_seguro(nome.upper().replace(" ", "_"), "COLUNA")
    candidato = base
    numero = 2
    while candidato in existentes:
        candidato = f"{base}_{numero}"
        numero += 1
    existentes.add(candidato)
    return candidato


def extrair_tabela(page: Page, credencial: Credencial, vigencia: str) -> list[dict[str, str]]:
    tabela = tabela_boletos(page)
    if not tabela.count():
        return []
    linhas = tabela.locator("tbody tr")
    if linhas.count() == 0:
        return []
    qtd_colunas = max(linhas.nth(i).locator("td").count() for i in range(linhas.count()))
    nomes_brutos = cabecalhos_fisicos(tabela, qtd_colunas)
    existentes: set[str] = set()
    nomes = [nome_coluna_unico(nome, existentes) for nome in nomes_brutos]
    registros: list[dict[str, str]] = []
    for indice in range(linhas.count()):
        linha = linhas.nth(indice)
        celulas = linha.locator("td")
        valores = [texto(celulas.nth(i).inner_text()) for i in range(celulas.count())]
        if not any(valores):
            continue
        registro = {
            "DATA_COLETA": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "VIGENCIA_SELECIONADA": vigencia,
            "LINHA_PLANILHA_CONTROLE": str(credencial.linha_excel),
            "USUARIO_MASCARADO": mascarar_usuario(credencial.usuario),
            "CONTRATO_CONTROLE": credencial.contrato,
        }
        for posicao, nome in enumerate(nomes):
            registro[nome] = valores[posicao] if posicao < len(valores) else ""
        registro["QTD_LINKS_RELATORIOS"] = str(celulas.last.locator("a").count()) if celulas.count() else "0"
        registros.append(registro)
    return registros


def linhas_presentes_na_tabela(page: Page) -> tuple[LinhaBoletoPortal, ...]:
    """Lê cada linha do portal pelas posições estáveis mostradas na tela."""
    tabela = tabela_boletos(page)
    if not tabela.count():
        return ()
    linhas = tabela.locator("tbody tr")
    if not linhas.count():
        return ()

    linhas_portal: list[LinhaBoletoPortal] = []
    chaves_lidas: set[str] = set()
    for indice in range(linhas.count()):
        celulas = linhas.nth(indice).locator("td")
        # Ordem da tabela: Ações, Código, Códigos Agrupados, Número,
        # Competência, Emissão, Vencimento, Valor, ...
        if celulas.count() < 8:
            continue
        codigo = texto(celulas.nth(1).inner_text())
        numero = texto(celulas.nth(3).inner_text())
        vencimento = texto(celulas.nth(6).inner_text())
        valor = texto(celulas.nth(7).inner_text())
        if not codigo or not numero:
            continue
        chave = apenas_digitos(numero) or f"{codigo}|{numero}"
        if chave in chaves_lidas:
            continue
        chaves_lidas.add(chave)
        linhas_portal.append(
            LinhaBoletoPortal(
                codigo=codigo,
                numero=numero,
                vencimento=vencimento,
                valor=valor,
            )
        )
    return tuple(linhas_portal)


def localizar_linha_boleto(page: Page, dados_linha: LinhaBoletoPortal) -> Locator | None:
    linhas = tabela_boletos(page).locator("tbody tr")
    for indice in range(linhas.count()):
        linha = linhas.nth(indice)
        celulas = linha.locator("td")
        if celulas.count() < 4:
            continue
        codigo = texto(celulas.nth(1).inner_text())
        numero = texto(celulas.nth(3).inner_text())
        if codigo == dados_linha.codigo and numero == dados_linha.numero:
            return linha
    return None


def fechar_modal_notas(page: Page) -> None:
    seletores = [
        "div.modal:visible button.close",
        "div.modal:visible [data-dismiss='modal']",
        "div.modal:visible .modal-header .close",
        "div[role='dialog']:visible button[aria-label='Close']",
    ]
    for seletor in seletores:
        try:
            loc = page.locator(seletor).first
            if loc.count() and loc.is_visible(timeout=500):
                loc.click()
                page.wait_for_timeout(500)
                return
        except Exception:
            pass
    page.keyboard.press("Escape")
    page.wait_for_timeout(500)


def extrair_notas_modal(page: Page, codigo: str) -> list[tuple[str, str, str]]:
    titulo = page.get_by_text(re.compile(r"Relação das Notas Existentes", re.I)).first
    titulo.wait_for(state="visible", timeout=TIMEOUT_PADRAO_MS)
    modal = page.locator("div.modal:visible, div[role='dialog']:visible").first
    if not modal.count():
        modal = titulo.locator("xpath=ancestor::div[contains(@class,'modal')][1]")
    linhas = modal.locator("table tbody tr")
    notas: list[tuple[str, str, str]] = []
    for indice in range(linhas.count()):
        celulas = linhas.nth(indice).locator("td")
        valores = [texto(celulas.nth(i).inner_text()) for i in range(celulas.count())]
        if len(valores) >= 2 and valores[0] and valores[1]:
            notas.append((valores[0], valores[1], valores[2] if len(valores) > 2 else ""))
    if not notas:
        raise RuntimeError(f"Código {codigo}: a janela abriu, mas nenhuma nota foi identificada.")
    return notas


def localizar_campos_prefeitura(page: Page) -> tuple[Locator, Locator, Locator]:
    expressoes = (
        "//*[contains(normalize-space(.), 'Inscrição Municipal Prestador')]/following::input[1]",
        "//*[contains(normalize-space(.), 'Número NFS-e')]/following::input[1]",
        "//*[contains(normalize-space(.), 'Código de autenticidade')]/following::input[1]",
    )
    encontrados: list[Locator] = []
    for expressao in expressoes:
        loc = page.locator(f"xpath={expressao}").first
        try:
            if loc.count() and loc.is_visible(timeout=1_000):
                encontrados.append(loc)
        except Exception:
            pass
    if len(encontrados) == 3:
        return encontrados[0], encontrados[1], encontrados[2]
    visiveis = page.locator("form input:visible")
    if visiveis.count() < 3:
        raise RuntimeError("Os três campos da consulta de autenticidade da Prefeitura não foram encontrados.")
    return visiveis.nth(0), visiveis.nth(1), visiveis.nth(2)


def salvar_download_com_nome(download: Download, destino: Path) -> tuple[Path, bool]:
    if destino.exists() and destino.stat().st_size > 0:
        try:
            download.cancel()
        except Exception:
            pass
        return destino, True
    destino.parent.mkdir(parents=True, exist_ok=True)
    download.save_as(destino)
    return destino, False


def aguardar_nova_janela(
    contexto: BrowserContext,
    paginas_antes: set[Page],
    timeout_segundos: int = 15,
) -> Page | None:
    limite = time.monotonic() + timeout_segundos
    while time.monotonic() < limite:
        novas = [pagina for pagina in contexto.pages if pagina not in paginas_antes]
        if novas:
            return novas[-1]
        time.sleep(0.25)
    return None


def bytes_sao_pdf(conteudo: bytes) -> bool:
    return len(conteudo) >= 500 and conteudo.lstrip().startswith(b"%PDF-")


def gravar_pdf_atomico(conteudo: bytes, destino: Path) -> Path:
    if not bytes_sao_pdf(conteudo):
        raise ValueError("O conteúdo recebido não possui uma assinatura PDF válida.")
    destino.parent.mkdir(parents=True, exist_ok=True)
    temporario = destino.with_name(f"{destino.name}.parcial")
    temporario.write_bytes(conteudo)
    if not bytes_sao_pdf(temporario.read_bytes()):
        temporario.unlink(missing_ok=True)
        raise ValueError("O arquivo temporário não passou na validação de PDF.")
    temporario.replace(destino)
    return destino


def tentar_pdf_de_respostas(
    respostas: list[Response],
    destino: Path,
    logger: logging.Logger,
) -> Path | None:
    logger.info("DIAGNOSTICO_PDF | ETAPA=RESPOSTAS | TOTAL_CAPTURADAS=%s", len(respostas))
    candidatos_analisados = 0
    for resposta in reversed(respostas):
        try:
            tipo = resposta.headers.get("content-type", "").lower()
            url_resposta = resposta.url.lower()
            if "pdf" not in tipo and not any(
                termo in url_resposta
                for termo in ("gerarrelatorio", "relatorio", "imprimir", "download")
            ):
                continue
            candidatos_analisados += 1
            logger.info(
                "DIAGNOSTICO_PDF | ETAPA=RESPOSTA_CANDIDATA | STATUS=%s | TIPO=%s | URL=%s",
                resposta.status,
                tipo or "SEM_CONTENT_TYPE",
                resposta.url,
            )
            conteudo = b""
            ultimo_erro = ""
            for tentativa in range(1, 4):
                try:
                    conteudo = resposta.body()
                    break
                except Exception as exc:
                    ultimo_erro = str(exc)
                    logger.warning(
                        "DIAGNOSTICO_PDF | ETAPA=LER_CORPO_RESPOSTA | TENTATIVA=%s/3 | ERRO=%s | URL=%s",
                        tentativa,
                        exc,
                        resposta.url,
                    )
                    time.sleep(1)
            logger.info(
                "DIAGNOSTICO_PDF | ETAPA=VALIDAR_RESPOSTA | BYTES=%s | ASSINATURA_HEX=%s | URL=%s",
                len(conteudo),
                conteudo[:12].hex() if conteudo else "VAZIO",
                resposta.url,
            )
            if bytes_sao_pdf(conteudo):
                logger.info("PDF capturado diretamente da resposta do portal: %s", resposta.url)
                return gravar_pdf_atomico(conteudo, destino)

            try:
                requisicao_original = resposta.request
                logger.info(
                    "DIAGNOSTICO_PDF | ETAPA=REPETIR_REQUISICAO | METODO=%s | URL=%s",
                    requisicao_original.method,
                    requisicao_original.url,
                )
                resposta_repetida = requisicao_original.frame.page.context.request.fetch(
                    requisicao_original,
                    timeout=TIMEOUT_DOWNLOAD_MS,
                )
                conteudo_repetido = resposta_repetida.body()
                logger.info(
                    "DIAGNOSTICO_PDF | ETAPA=VALIDAR_REQUISICAO_REPETIDA | STATUS=%s | TIPO=%s | BYTES=%s | ASSINATURA_HEX=%s | URL=%s",
                    resposta_repetida.status,
                    resposta_repetida.headers.get("content-type", "SEM_CONTENT_TYPE"),
                    len(conteudo_repetido),
                    conteudo_repetido[:12].hex() if conteudo_repetido else "VAZIO",
                    resposta_repetida.url,
                )
                if resposta_repetida.ok and bytes_sao_pdf(conteudo_repetido):
                    logger.info("PDF obtido ao repetir a requisição original do portal.")
                    return gravar_pdf_atomico(conteudo_repetido, destino)
                logger.error(
                    "DIAGNOSTICO_PDF | ETAPA=REPETIR_REQUISICAO | ERRO=CONTEUDO_NAO_E_PDF | URL=%s",
                    requisicao_original.url,
                )
            except Exception as exc_repeticao:
                logger.error(
                    "DIAGNOSTICO_PDF | ETAPA=REPETIR_REQUISICAO | ERRO=%s | URL=%s",
                    exc_repeticao,
                    resposta.url,
                )

            if ultimo_erro and not conteudo:
                logger.error(
                    "DIAGNOSTICO_PDF | ETAPA=RESPOSTA_REJEITADA | MOTIVO=%s | URL=%s",
                    ultimo_erro,
                    resposta.url,
                )
            else:
                logger.error(
                    "DIAGNOSTICO_PDF | ETAPA=RESPOSTA_REJEITADA | MOTIVO=CONTEUDO_NAO_E_PDF | URL=%s",
                    resposta.url,
                )
        except Exception as exc:
            logger.error(
                "DIAGNOSTICO_PDF | ETAPA=ANALISAR_RESPOSTA | ERRO=%s",
                exc,
            )
            continue
    if candidatos_analisados == 0:
        logger.error(
            "DIAGNOSTICO_PDF | ETAPA=RESPOSTAS | ERRO=NENHUMA_RESPOSTA_CANDIDATA_ENCONTRADA"
        )
    return None


def tentar_pdf_por_url(
    pagina: Page,
    url: str,
    destino: Path,
    logger: logging.Logger,
) -> Path | None:
    if not url or url.startswith(("about:", "edge:", "chrome:")):
        logger.error(
            "DIAGNOSTICO_PDF | ETAPA=URL_DIRETA | ERRO=URL_NAO_REQUISITAVEL | URL=%s",
            url or "VAZIA",
        )
        return None
    try:
        resposta = pagina.context.request.get(url, timeout=TIMEOUT_DOWNLOAD_MS)
        logger.info(
            "DIAGNOSTICO_PDF | ETAPA=URL_DIRETA | STATUS=%s | TIPO=%s | URL=%s",
            resposta.status,
            resposta.headers.get("content-type", "SEM_CONTENT_TYPE"),
            url,
        )
        if not resposta.ok:
            logger.error(
                "DIAGNOSTICO_PDF | ETAPA=URL_DIRETA | ERRO=HTTP_%s | URL=%s",
                resposta.status,
                url,
            )
            return None
        conteudo = resposta.body()
        logger.info(
            "DIAGNOSTICO_PDF | ETAPA=VALIDAR_URL | BYTES=%s | ASSINATURA_HEX=%s | URL=%s",
            len(conteudo),
            conteudo[:12].hex() if conteudo else "VAZIO",
            url,
        )
        if bytes_sao_pdf(conteudo):
            logger.info("PDF obtido diretamente pela URL autenticada: %s", url)
            return gravar_pdf_atomico(conteudo, destino)
        logger.error(
            "DIAGNOSTICO_PDF | ETAPA=URL_DIRETA | ERRO=CONTEUDO_NAO_E_PDF | URL=%s",
            url,
        )
    except Exception as exc:
        logger.error(
            "DIAGNOSTICO_PDF | ETAPA=URL_DIRETA | ERRO=%s | URL=%s",
            exc,
            url,
        )
        return None
    return None


def fontes_pdf_da_pagina(pagina: Page) -> list[str]:
    try:
        fontes = pagina.evaluate(
            """
            () => Array.from(document.querySelectorAll('iframe, embed, object'))
                .map(el => el.src || el.data || el.getAttribute('src') || el.getAttribute('data') || '')
                .filter(Boolean)
            """
        )
        return [str(fonte) for fonte in fontes if fonte]
    except Exception:
        return []


def obter_bytes_blob(
    pagina: Page,
    url_blob: str,
    logger: logging.Logger,
) -> bytes | None:
    try:
        base64_conteudo = pagina.evaluate(
            """
            async (url) => {
                let blob = null;
                if (window.__roboBlobsCriados && window.__roboBlobsCriados.has(url)) {
                    blob = window.__roboBlobsCriados.get(url);
                } else if (window.__roboBlobsCriados && window.__roboBlobsCriados.size) {
                    const blobs = Array.from(window.__roboBlobsCriados.values());
                    blob = blobs[blobs.length - 1];
                } else {
                    const resposta = await fetch(url);
                    blob = await resposta.blob();
                }
                const buffer = await blob.arrayBuffer();
                const bytes = new Uint8Array(buffer);
                let binario = '';
                const bloco = 0x8000;
                for (let i = 0; i < bytes.length; i += bloco) {
                    binario += String.fromCharCode(...bytes.subarray(i, i + bloco));
                }
                return btoa(binario);
            }
            """,
            url_blob,
        )
        conteudo = base64.b64decode(base64_conteudo)
        logger.info(
            "DIAGNOSTICO_BLOB | STATUS=CAPTURADO | URL=%s | BYTES=%s | ASSINATURA_HEX=%s | PAGINA_ORIGEM=%s",
            url_blob,
            len(conteudo),
            conteudo[:12].hex() if conteudo else "VAZIO",
            pagina.url,
        )
        return conteudo
    except Exception as exc:
        logger.error(
            "DIAGNOSTICO_BLOB | STATUS=ERRO | ERRO=%s | URL=%s | PAGINA_ORIGEM=%s",
            exc,
            url_blob,
            pagina.url,
        )
        return None


def tentar_pdf_blob(
    pagina: Page,
    url_blob: str,
    destino: Path,
    logger: logging.Logger,
) -> Path | None:
    conteudo = obter_bytes_blob(pagina, url_blob, logger)
    if conteudo is not None and bytes_sao_pdf(conteudo):
        logger.info("PDF recuperado diretamente do Blob criado pelo portal.")
        return gravar_pdf_atomico(conteudo, destino)
    if conteudo is not None:
        logger.error(
            "DIAGNOSTICO_PDF | ETAPA=BLOB | ERRO=CONTEUDO_NAO_E_PDF | BYTES=%s | URL=%s",
            len(conteudo),
            url_blob,
        )
    return None


def tentar_pdf_blob_em_paginas(
    paginas: list[Page],
    url_blob: str,
    destino: Path,
    logger: logging.Logger,
) -> Path | None:
    paginas_unicas: list[Page] = []
    for pagina in paginas:
        if pagina not in paginas_unicas:
            paginas_unicas.append(pagina)
    for pagina in paginas_unicas:
        caminho = tentar_pdf_blob(pagina, url_blob, destino, logger)
        if caminho is not None:
            return caminho
    return None


def salvar_pagina_sem_impressora(
    pagina_resultado: Page,
    destino: Path,
    logger: logging.Logger,
    respostas_capturadas: list[Response] | None = None,
    permitir_pdf_da_pagina: bool = True,
    aceitar_arquivo_existente: bool = True,
    pagina_origem: Page | None = None,
) -> tuple[Path, bool]:
    if aceitar_arquivo_existente and destino.exists() and destino.stat().st_size > 0:
        logger.info("PDF já existente: %s", destino.name)
        return destino, True
    if not aceitar_arquivo_existente and destino.exists():
        logger.warning(
            "DIAGNOSTICO_PDF | ETAPA=ARQUIVO_EXISTENTE | ACAO=IGNORAR_E_SUBSTITUIR_SOMENTE_SE_CAPTURAR_ORIGINAL | ARQUIVO=%s",
            destino,
        )

    destino.parent.mkdir(parents=True, exist_ok=True)
    aguardar_carregamento(pagina_resultado, 2_000)
    logger.info(
        "DIAGNOSTICO_PDF | ETAPA=INICIO | URL_PAGINA=%s | DESTINO=%s | PERMITIR_PDF_DA_PAGINA=%s",
        pagina_resultado.url,
        destino,
        permitir_pdf_da_pagina,
    )

    caminho = tentar_pdf_de_respostas(respostas_capturadas or [], destino, logger)
    if caminho is not None:
        return caminho, False

    if pagina_resultado.url.startswith("blob:"):
        paginas_blob = [pagina_resultado]
        if pagina_origem is not None:
            paginas_blob.append(pagina_origem)
        caminho = tentar_pdf_blob_em_paginas(
            paginas_blob,
            pagina_resultado.url,
            destino,
            logger,
        )
    else:
        caminho = tentar_pdf_por_url(pagina_resultado, pagina_resultado.url, destino, logger)
    if caminho is not None:
        return caminho, False

    fontes = fontes_pdf_da_pagina(pagina_resultado)
    logger.info(
        "DIAGNOSTICO_PDF | ETAPA=FONTES_HTML | TOTAL=%s | FONTES=%s",
        len(fontes),
        " | ".join(fontes[:10]) if fontes else "NENHUMA",
    )
    for fonte in fontes:
        if fonte.startswith("blob:"):
            paginas_blob = [pagina_resultado]
            if pagina_origem is not None:
                paginas_blob.append(pagina_origem)
            caminho = tentar_pdf_blob_em_paginas(paginas_blob, fonte, destino, logger)
        else:
            caminho = tentar_pdf_por_url(
                pagina_resultado,
                urljoin(pagina_resultado.url, fonte),
                destino,
                logger,
            )
        if caminho is not None:
            return caminho, False

    if permitir_pdf_da_pagina:
        try:
            # Força o navegador a usar o estilo visual da tela
            pagina_resultado.emulate_media(media="screen")
            
            # Injeta CSS para forçar a quebra de texto, remover barras de rolagem e ajustar larguras
            pagina_resultado.add_style_tag(content="""
                * {
                    overflow: visible !important;
                    white-space: normal !important;
                    word-wrap: break-word !important;
                }
                body, html, form, table {
                    width: 100% !important;
                    max-width: 100% !important;
                    min-width: 0 !important;
                    margin: 0 auto !important;
                }
            """)
            
            # Aguarda a aplicação do CSS e renderização
            pagina_resultado.wait_for_timeout(1000)

            conteudo_gerado = pagina_resultado.pdf(
                format="A4",
                print_background=True,
                margin={"top": "10mm", "right": "10mm", "bottom": "10mm", "left": "10mm"},
                scale=0.70  # Escala reduzida para 70% para garantir que tabelas largas caibam na folha
            )
            
            if bytes_sao_pdf(conteudo_gerado):
                logger.warning(
                    "DIAGNOSTICO_PDF | ETAPA=PDF_DA_PAGINA | AVISO=PDF_GERADO_A_PARTIR_DA_PAGINA_HTML"
                )
                return gravar_pdf_atomico(conteudo_gerado, destino), False
        except Exception as exc:
            logger.error(
                "DIAGNOSTICO_PDF | ETAPA=PDF_DA_PAGINA | ERRO=%s",
                exc,
            )
            if bytes_sao_pdf(conteudo_gerado):
                logger.warning(
                    "DIAGNOSTICO_PDF | ETAPA=PDF_DA_PAGINA | AVISO=PDF_GERADO_A_PARTIR_DA_PAGINA_HTML"
                )
                return gravar_pdf_atomico(conteudo_gerado, destino), False
        except Exception as exc:
            logger.error(
                "DIAGNOSTICO_PDF | ETAPA=PDF_DA_PAGINA | ERRO=%s",
                exc,
            )
    else:
        logger.error(
            "DIAGNOSTICO_PDF | ETAPA=PDF_DA_PAGINA | BLOQUEADO=SIM | MOTIVO=EVITAR_PDF_COM_APARENCIA_DE_PRINT"
        )

    logger.error(
        "DIAGNOSTICO_PDF | ETAPA=FIM | STATUS=ERRO | URL_PAGINA=%s | DESTINO=%s",
        pagina_resultado.url,
        destino,
    )
    raise RuntimeError(
        "Não foi possível capturar o PDF original com segurança. Consulte no log as linhas "
        "iniciadas por DIAGNOSTICO_PDF. Nenhum PDF com aparência de print foi criado e "
        "nenhum comando foi enviado à impressora."
    )


def conteudo_binario_valido(conteudo: bytes, extensao: str) -> bool:
    if not conteudo:
        return False
    if extensao.lower() == ".pdf":
        return bytes_sao_pdf(conteudo)
    inicio = conteudo.lstrip()[:100].lower()
    return not inicio.startswith((b"<!doctype html", b"<html", b"<head", b"<body"))


def gravar_arquivo_atomico(conteudo: bytes, destino: Path) -> Path:
    if not conteudo_binario_valido(conteudo, destino.suffix):
        raise ValueError(f"O conteúdo recebido não é um arquivo {destino.suffix or 'válido'}.")
    destino.parent.mkdir(parents=True, exist_ok=True)
    temporario = destino.with_name(f"{destino.name}.parcial")
    temporario.write_bytes(conteudo)
    temporario.replace(destino)
    return destino


def renomear_arquivo_baixado(caminho: Path, nome_final: str, logger: logging.Logger) -> tuple[Path, bool]:
    """Renomeia somente um arquivo já salvo e validado, sem alterar o download."""
    destino = caminho.with_name(nome_final)
    if caminho == destino:
        return caminho, False
    if destino.exists() and destino.stat().st_size > 0:
        logger.info("Arquivo final já existente; mantendo: %s", destino.name)
        return destino, True
    caminho.replace(destino)
    logger.info("Arquivo renomeado: %s -> %s", caminho.name, destino.name)
    return destino, False


def ultimo_blob_criado(pagina: Page) -> str:
    try:
        return texto(
            pagina.evaluate(
                """
                () => {
                    if (!window.__roboBlobsCriados || !window.__roboBlobsCriados.size) return '';
                    const urls = Array.from(window.__roboBlobsCriados.keys());
                    return urls[urls.length - 1] || '';
                }
                """
            )
        )
    except Exception:
        return ""


def salvar_arquivo_nova_aba(
    pagina_resultado: Page,
    pagina_origem: Page,
    destino: Path,
    logger: logging.Logger,
    respostas_capturadas: list[Response],
) -> tuple[Path, bool]:
    limite_url = time.monotonic() + 15
    while time.monotonic() < limite_url and pagina_resultado.url in {"", "about:blank"}:
        pagina_resultado.wait_for_timeout(250)
    logger.info(
        "ETAPA_ROBO | STATUS=DETALHE | ETAPA=NOVA_ABA_DETECTADA | URL=%s | DESTINO=%s",
        pagina_resultado.url,
        destino,
    )
    if destino.suffix.lower() == ".pdf":
        return salvar_pagina_sem_impressora(
            pagina_resultado,
            destino,
            logger,
            respostas_capturadas,
            permitir_pdf_da_pagina=False,
            pagina_origem=pagina_origem,
        )

    if destino.exists() and destino.stat().st_size > 0:
        return destino, True
    url_blob = pagina_resultado.url if pagina_resultado.url.startswith("blob:") else ultimo_blob_criado(pagina_origem)
    if url_blob:
        for pagina_blob in (pagina_origem, pagina_resultado):
            conteudo = obter_bytes_blob(pagina_blob, url_blob, logger)
            if conteudo_binario_valido(conteudo or b"", destino.suffix):
                return gravar_arquivo_atomico(conteudo or b"", destino), False
    for resposta in reversed(respostas_capturadas):
        try:
            tipo = resposta.headers.get("content-type", "").lower()
            if not any(termo in tipo for termo in ("csv", "excel", "octet-stream")):
                continue
            conteudo = resposta.body()
            if conteudo_binario_valido(conteudo, destino.suffix):
                return gravar_arquivo_atomico(conteudo, destino), False
        except Exception as exc:
            logger.warning("Falha ao analisar resposta do arquivo %s: %s", destino.name, exc)
    raise RuntimeError(f"A nova aba foi detectada, mas o conteúdo de {destino.name} não pôde ser capturado.")


def clicar_e_salvar_resultado(
    page: Page,
    elemento: Locator,
    destino: Path,
    codigo: str,
    descricao: str,
    logger: logging.Logger,
) -> tuple[Path, bool]:
    downloads: list[Download] = []
    paginas_evento: list[Page] = []
    respostas: list[Response] = []
    paginas_antes = set(page.context.pages)

    def ao_download(download: Download) -> None:
        downloads.append(download)

    def ao_pagina(nova_pagina: Page) -> None:
        paginas_evento.append(nova_pagina)

    def ao_resposta(resposta: Response) -> None:
        respostas.append(resposta)

    page.on("download", ao_download)
    page.on("popup", ao_pagina)
    page.context.on("page", ao_pagina)
    page.context.on("response", ao_resposta)
    try:
        try:
            page.evaluate(
                "() => { if (window.__roboBlobsCriados) window.__roboBlobsCriados.clear(); }"
            )
        except Exception as exc:
            logger.warning("Não foi possível limpar o histórico de Blobs antes de '%s': %s", descricao, exc)
        elemento.wait_for(state="visible", timeout=TIMEOUT_PADRAO_MS)
        elemento.click()
        limite = time.monotonic() + (TIMEOUT_DOWNLOAD_MS / 1_000)
        while time.monotonic() < limite:
            if downloads:
                download = downloads[0]
                sufixo_real = Path(download.suggested_filename).suffix
                destino_real = destino.with_suffix(sufixo_real) if sufixo_real else destino
                return salvar_download_com_nome(download, destino_real)

            novas_por_lista = [pagina for pagina in page.context.pages if pagina not in paginas_antes]
            candidatas = paginas_evento + novas_por_lista
            pagina_secundaria = next(
                (pagina for pagina in reversed(candidatas) if pagina is not page and not pagina.is_closed()),
                None,
            )
            if pagina_secundaria is not None:
                logger.info(
                    "Código %s: nova aba detectada para '%s'.",
                    codigo,
                    descricao,
                )
                return salvar_arquivo_nova_aba(
                    pagina_secundaria,
                    page,
                    destino,
                    logger,
                    respostas,
                )

            url_blob = ultimo_blob_criado(page)
            if url_blob:
                conteudo = obter_bytes_blob(page, url_blob, logger)
                if conteudo_binario_valido(conteudo or b"", destino.suffix):
                    return gravar_arquivo_atomico(conteudo or b"", destino), False
            time.sleep(0.25)
        raise TimeoutError(
            f"O portal não disparou download, popup nem Blob válido em {TIMEOUT_DOWNLOAD_MS // 1_000} segundos."
        )
    finally:
        for emissor, evento, funcao in (
            (page, "download", ao_download),
            (page, "popup", ao_pagina),
            (page.context, "page", ao_pagina),
            (page.context, "response", ao_resposta),
        ):
            try:
                emissor.remove_listener(evento, funcao)
            except Exception:
                pass
        for pagina_aberta in list(page.context.pages):
            if pagina_aberta not in paginas_antes and pagina_aberta is not page:
                try:
                    pagina_aberta.close()
                except Exception:
                    pass


def baixar_nf_prefeitura(
    contexto: BrowserContext,
    inscricao: str,
    codigo_portal: str,
    numero_nota: str,
    codigo_verificador: str,
    pasta_destino: Path,
    logger: logging.Logger,
) -> tuple[Path, bool]:
    page = contexto.new_page()
    respostas_capturadas: list[Response] = []

    def registrar_resposta(resposta: Response) -> None:
        respostas_capturadas.append(resposta)

    try:
        contexto_log = {"codigo": codigo_portal, "nota": numero_nota}
        with etapa_registrada(logger, "PREFEITURA_ABRIR_CONSULTA", **contexto_log):
            page.goto(URL_PREFEITURA, wait_until="domcontentloaded", timeout=TIMEOUT_CARREGAMENTO_MS)
            aguardar_carregamento(page, 500)
        with etapa_registrada(logger, "PREFEITURA_LOCALIZAR_CAMPOS", **contexto_log):
            campo_inscricao, campo_numero, campo_autenticidade = localizar_campos_prefeitura(page)
        with etapa_registrada(logger, "PREFEITURA_PREENCHER_CONSULTA", **contexto_log):
            campo_inscricao.fill(inscricao)
            campo_numero.fill(numero_nota)
            campo_autenticidade.fill(codigo_verificador)

        paginas_antes = set(contexto.pages)
        with etapa_registrada(logger, "PREFEITURA_ENVIAR_CONSULTA", **contexto_log):
            verificar = page.locator(f"xpath={XPATH_VERIFICAR_PREFEITURA}")
            verificar.wait_for(state="visible", timeout=TIMEOUT_PADRAO_MS)
            contexto.on("response", registrar_resposta)
            verificar.click()
        with etapa_registrada(logger, "PREFEITURA_IDENTIFICAR_PAGINA_RESULTADO", **contexto_log):
            pagina_secundaria = aguardar_nova_janela(contexto, paginas_antes)
        pagina_resultado = pagina_secundaria if pagina_secundaria is not None else page
        if pagina_secundaria is not None:
            logger.info("A nota fiscal abriu em uma segunda janela; capturando o PDF sem usar impressão.")
        with etapa_registrada(
            logger,
            "PREFEITURA_AGUARDAR_RESULTADO",
            **contexto_log,
            nova_janela="SIM" if pagina_secundaria is not None else "NAO",
        ):
            aguardar_carregamento(pagina_resultado, 2_000)

        destino = pasta_destino / f"NFSe_{nome_seguro(codigo_portal)}_{nome_seguro(numero_nota)}.pdf"
        ultimo_erro = ""
        for tentativa in range(1, MAXIMO_TENTATIVAS_DOWNLOAD + 1):
            try:
                with etapa_registrada(
                    logger,
                    "PREFEITURA_CAPTURAR_PDF_ORIGINAL",
                    **contexto_log,
                    tentativa=f"{tentativa}/{MAXIMO_TENTATIVAS_DOWNLOAD}",
                    url=pagina_resultado.url,
                ):
                    caminho, existente = salvar_pagina_sem_impressora(
                        pagina_resultado,
                        destino,
                        logger,
                        respostas_capturadas,
                        permitir_pdf_da_pagina=True,  # CORRIGIDO: Agora permite usar .pdf() caso falhe
                        aceitar_arquivo_existente=False,
                    )
                logger.info(
                    "NFS-e %s do código %s: %s",
                    numero_nota,
                    codigo_portal,
                    "já existente" if existente else "salva como PDF",
                )
                return caminho, existente
            except Exception as exc:
                ultimo_erro = str(exc)
                logger.warning(
                    "NFS-e %s: tentativa %s/%s de salvamento seguro falhou: %s",
                    numero_nota,
                    tentativa,
                    MAXIMO_TENTATIVAS_DOWNLOAD,
                    exc,
                )
                pagina_resultado.wait_for_timeout(1_000)
        raise RuntimeError(f"A NFS-e {numero_nota} não foi salva como PDF: {ultimo_erro}")
    finally:
        try:
            contexto.remove_listener("response", registrar_resposta)
        except Exception:
            pass
        for pagina in list(contexto.pages):
            if pagina is not page and "camacari.ba.gov.br" in pagina.url:
                try:
                    pagina.close()
                except Exception:
                    pass
        try:
            page.close()
        except Exception:
            pass


def baixar_notas_codigo(
    page: Page,
    contexto: BrowserContext,
    linha: Locator,
    credencial: Credencial,
    codigo: str,
    pasta_codigo: Path,
    logger: logging.Logger,
    dados: DadosExecucao,
) -> tuple[int, int, int]:
    with etapa_registrada(logger, "PORTAL_LOCALIZAR_RELACAO_NOTAS", codigo=codigo):
        relatorios = linha.locator("td").last.locator("a")
    if relatorios.count() < 1:
        logger.warning("Código %s: Relação das Notas Existentes indisponível.", codigo)
        return 0, 0, 1
    with etapa_registrada(logger, "PORTAL_ABRIR_RELACAO_NOTAS", codigo=codigo):
        relatorios.nth(0).click()
    with etapa_registrada(logger, "PORTAL_EXTRAIR_NOTAS", codigo=codigo):
        notas = extrair_notas_modal(page, codigo)
    baixados = ignorados = falhas = 0
    for numero, verificador, cpf_cnpj in notas:
        registro = {
            "DATA_COLETA": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "LINHA_PLANILHA_CONTROLE": str(credencial.linha_excel),
            "CONTRATO": credencial.contrato,
            "CODIGO_PORTAL": codigo,
            "NUMERO_NOTA": numero,
            "CODIGO_VERIFICADOR": verificador,
            "CPF_CNPJ_PRESTADOR": cpf_cnpj,
            "INSCRICAO_MUNICIPAL_PRESTADOR": credencial.inscricao_municipal,
            "STATUS_DOWNLOAD": "PENDENTE",
            "ARQUIVO": "",
            "DETALHE": "",
        }
        try:
            caminho, existente = baixar_nf_prefeitura(
                contexto,
                credencial.inscricao_municipal,
                codigo,
                numero,
                verificador,
                pasta_codigo,
                logger,
            )
            caminho, nome_final_existente = renomear_arquivo_baixado(
                caminho,
                f"{nome_seguro(credencial.contrato)}_NFS_{nome_seguro(numero)}_{nome_seguro(codigo)}.pdf",
                logger,
            )
            existente = existente or nome_final_existente
            registro["STATUS_DOWNLOAD"] = "JA_EXISTENTE" if existente else "BAIXADO"
            registro["ARQUIVO"] = caminho.name
            if existente:
                ignorados += 1
            else:
                baixados += 1
        except Exception as exc:
            falhas += 1
            registro["STATUS_DOWNLOAD"] = "ERRO"
            registro["DETALHE"] = str(exc)
            logger.exception(
                "ETAPA_ROBO | STATUS=ERRO | ETAPA=PROCESSAR_NFSE | CODIGO=%s | NOTA=%s | TIPO_ERRO=%s | ERRO=%s",
                codigo,
                numero,
                type(exc).__name__,
                exc,
            )
        dados.registros_notas.append(registro)
    with etapa_registrada(logger, "PORTAL_FECHAR_RELACAO_NOTAS", codigo=codigo):
        fechar_modal_notas(page)
    return baixados, ignorados, falhas


def baixar_relatorio_icone(
    page: Page,
    linha: Locator,
    tipo_relatorio: str,
    codigo: str,
    descricao: str,
    extensao: str,
    contrato: str,
    pasta_codigo: Path,
    logger: logging.Logger,
) -> tuple[str, bool]:
    with etapa_registrada(
        logger,
        "PORTAL_LOCALIZAR_RELATORIO",
        codigo=codigo,
        relatorio=descricao,
    ):
        relatorios = linha.locator("td").last.locator("a")
        indice_link = localizar_indice_relatorio(relatorios, tipo_relatorio)
    if indice_link is None:
        logger.warning("Código %s: relatório '%s' indisponível.", codigo, descricao)
        return "INDISPONIVEL", False
    extensao_final = ".pdf" if tipo_relatorio == "coparticipacao" else extensao
    destino_esperado = (
        pasta_codigo
        / f"{nome_seguro(contrato)}_Coparticipacao_{nome_seguro(codigo)}{extensao_final}"
    )
    if destino_esperado.exists() and destino_esperado.stat().st_size > 0:
        logger.info("Código %s: '%s' já existente: %s", codigo, descricao, destino_esperado.name)
        return "JA_EXISTENTE", True

    ultimo_erro = ""
    for tentativa in range(1, MAXIMO_TENTATIVAS_DOWNLOAD + 1):
        try:
            with etapa_registrada(
                logger,
                "PORTAL_DISPARAR_RELATORIO",
                codigo=codigo,
                relatorio=descricao,
                tentativa=f"{tentativa}/{MAXIMO_TENTATIVAS_DOWNLOAD}",
            ):
                relatorios = linha.locator("td").last.locator("a")
                caminho, existente = clicar_e_salvar_resultado(
                    page,
                    relatorios.nth(indice_link),
                    destino_esperado,
                    codigo,
                    descricao,
                    logger,
                )
            logger.info(
                "Código %s: '%s' %s.",
                codigo,
                descricao,
                "já existente" if existente else f"salvo em {caminho.name}",
            )
            return "JA_EXISTENTE" if existente else "BAIXADO", existente
        except Exception as exc:
            ultimo_erro = str(exc)
            logger.exception(
                "ETAPA_ROBO | STATUS=ERRO | ETAPA=BAIXAR_RELATORIO | CODIGO=%s | RELATORIO=%s | TENTATIVA=%s/%s | TIPO_ERRO=%s | ERRO=%s",
                codigo,
                descricao,
                tentativa,
                MAXIMO_TENTATIVAS_DOWNLOAD,
                type(exc).__name__,
                exc,
            )
            page.wait_for_timeout(1_000)
    raise RuntimeError(f"Código {codigo}, relatório '{descricao}': {ultimo_erro}")


def baixar_boletas_modal(
    page: Page,
    linha: Locator,
    codigo: str,
    contrato: str,
    pasta_codigo: Path,
    logger: logging.Logger,
) -> tuple[int, int, int, int]:
    """Baixa separadamente CSV e PDF do modal Relação das Boletas Existentes."""
    relatorios = linha.locator("td").last.locator("a")
    indice_link = localizar_indice_relatorio(relatorios, "boletas")
    if indice_link is None:
        logger.warning("Código %s: Relação das Boletas Existentes indisponível.", codigo)
        return 0, 0, 2, 0

    arquivos = (
        (
            "CSV",
            XPATH_BOLETAS_MODAL_CSV,
            pasta_codigo
            / f"{nome_seguro(contrato)}_Relacao_das_Boletas_Existentes_{nome_seguro(codigo)}.csv",
        ),
        (
            "PDF",
            XPATH_BOLETAS_MODAL_PDF,
            pasta_codigo
            / f"{nome_seguro(contrato)}_Relacao_das_Boletas_Existentes_{nome_seguro(codigo)}.pdf",
        ),
    )
    baixados = ignorados = indisponiveis = falhas = 0
    try:
        for formato, xpath_botao, destino in arquivos:
            if destino.exists() and destino.stat().st_size > 0:
                ignorados += 1
                logger.info("Código %s: Boletas %s já existente: %s", codigo, formato, destino.name)
                continue

            ultimo_erro = ""
            concluido = False
            for tentativa in range(1, MAXIMO_TENTATIVAS_DOWNLOAD + 1):
                try:
                    botao = page.locator(f"xpath={xpath_botao}")
                    if not botao.is_visible(timeout=1_000):
                        with etapa_registrada(
                            logger,
                            "PORTAL_ABRIR_MODAL_BOLETAS",
                            codigo=codigo,
                            formato=formato,
                        ):
                            relatorios = linha.locator("td").last.locator("a")
                            relatorios.nth(indice_link).click()
                            botao.wait_for(state="visible", timeout=TIMEOUT_PADRAO_MS)

                    with etapa_registrada(
                        logger,
                        "PORTAL_BAIXAR_BOLETAS_MODAL",
                        codigo=codigo,
                        formato=formato,
                        tentativa=f"{tentativa}/{MAXIMO_TENTATIVAS_DOWNLOAD}",
                    ):
                        caminho, existente = clicar_e_salvar_resultado(
                            page,
                            botao,
                            destino,
                            codigo,
                            f"Relação das Boletas Existentes {formato}",
                            logger,
                        )
                    if existente:
                        ignorados += 1
                    else:
                        baixados += 1
                    logger.info("Código %s: Boletas %s salva em %s.", codigo, formato, caminho.name)
                    concluido = True
                    break
                except Exception as exc:
                    ultimo_erro = str(exc)
                    logger.exception(
                        "ETAPA_ROBO | STATUS=ERRO | ETAPA=BAIXAR_BOLETAS_MODAL | CODIGO=%s | FORMATO=%s | TENTATIVA=%s/%s | TIPO_ERRO=%s | ERRO=%s",
                        codigo,
                        formato,
                        tentativa,
                        MAXIMO_TENTATIVAS_DOWNLOAD,
                        type(exc).__name__,
                        exc,
                    )
                    page.wait_for_timeout(1_000)
            if not concluido:
                falhas += 1
                logger.error("Código %s: não foi possível salvar Boletas %s: %s", codigo, formato, ultimo_erro)
    finally:
        try:
            fechar_modal_notas(page)
        except Exception:
            pass
    return baixados, ignorados, indisponiveis, falhas


def baixar_boleto_codigo(
    page: Page,
    linha: Locator,
    codigo: str,
    numero: str,
    vencimento: str,
    valor: str,
    contrato: str,
    pasta_codigo: Path,
    logger: logging.Logger,
) -> tuple[str, bool]:
    """Baixa o boleto e usa os dados da própria linha para nomear o arquivo."""
    destino = pasta_codigo / (
        f"{nome_seguro(contrato)}_Boleto_Venc_{data_para_nome(vencimento)}_"
        f"Valor_{valor_para_nome(valor)}_{nome_seguro(codigo)}.pdf"
    )
    if destino.exists() and destino.stat().st_size > 0:
        logger.info("Código %s | Número %s: boleto já existente: %s", codigo, numero, destino.name)
        return "JA_EXISTENTE", True

    ultimo_erro = ""
    for tentativa in range(1, MAXIMO_TENTATIVAS_DOWNLOAD + 1):
        try:
            with etapa_registrada(
                logger,
                "PORTAL_LOCALIZAR_BOTAO_BOLETO",
                codigo=codigo,
                numero=numero,
                tentativa=f"{tentativa}/{MAXIMO_TENTATIVAS_DOWNLOAD}",
            ):
                # Equivale aos XPaths absolutos informados, por exemplo:
                # /html/body/main/div[3]/div[3]/div[1]/table/tbody/tr[1]/td[1]/button[2]
                # /html/body/main/div[3]/div[3]/div[1]/table/tbody/tr[4]/td[1]/button[2]
                botao = linha.locator("xpath=./td[1]/button[2]")
                if not botao.count():
                    logger.warning("Código %s: botão de download do boleto indisponível.", codigo)
                    return "INDISPONIVEL", False

            with etapa_registrada(
                logger,
                "PORTAL_BAIXAR_BOLETO",
                codigo=codigo,
                numero=numero,
                tentativa=f"{tentativa}/{MAXIMO_TENTATIVAS_DOWNLOAD}",
            ):
                caminho, existente = clicar_e_salvar_resultado(
                    page,
                    botao,
                    destino,
                    codigo,
                    "Boleto",
                    logger,
                )
            logger.info(
                "Código %s | Número %s: boleto %s.",
                codigo,
                numero,
                "já existente" if existente else f"salvo em {caminho.name}",
            )
            return "JA_EXISTENTE" if existente else "BAIXADO", existente
        except Exception as exc:
            ultimo_erro = str(exc)
            logger.exception(
                "ETAPA_ROBO | STATUS=ERRO | ETAPA=BAIXAR_BOLETO | CODIGO=%s | NUMERO=%s | TENTATIVA=%s/%s | TIPO_ERRO=%s | ERRO=%s",
                codigo,
                numero,
                tentativa,
                MAXIMO_TENTATIVAS_DOWNLOAD,
                type(exc).__name__,
                exc,
            )
            page.wait_for_timeout(1_000)
    raise RuntimeError(f"Código {codigo}, boleto: {ultimo_erro}")


def localizar_indice_relatorio(relatorios: Locator, tipo_relatorio: str) -> int | None:
    quantidade = relatorios.count()
    for indice in range(quantidade):
        link = relatorios.nth(indice)
        partes = [
            texto(link.get_attribute("title")),
            texto(link.get_attribute("data-original-title")),
            texto(link.get_attribute("aria-label")),
            texto(link.get_attribute("href")),
            texto(link.get_attribute("onclick")),
            texto(link.inner_text()),
        ]
        try:
            partes.append(texto(link.locator("i, span").first.get_attribute("class")))
        except Exception:
            pass
        descricao_link = normalizar(" ".join(partes))
        if tipo_relatorio == "coparticipacao_csv":
            if "copart" in descricao_link and "csv" in descricao_link:
                return indice
        elif tipo_relatorio == "coparticipacao":
            if "copart" in descricao_link and "csv" not in descricao_link:
                return indice
        elif tipo_relatorio == "boletas":
            if any(termo in descricao_link for termo in ("boleta", "boleto", "relacao_boleta")):
                return indice

    # Fallback conforme a posição informada no portal. Quando existem somente
    # dois ícones, o print mostra Notas no primeiro e Boletas no último.
    if tipo_relatorio == "boletas" and quantidade >= 2:
        return quantidade - 1
    if tipo_relatorio == "coparticipacao" and quantidade >= 3:
        return 1
    if tipo_relatorio == "coparticipacao_csv" and quantidade >= 4:
        return 2
    return None


def salvar_evidencia(page: Page | None, pasta_logs: Path, linha_excel: int) -> None:
    if page is None:
        return
    try:
        pasta = pasta_logs / "Evidencias"
        pasta.mkdir(parents=True, exist_ok=True)
        page.screenshot(
            path=pasta / f"falha_linha_{linha_excel}_{datetime.now():%Y%m%d_%H%M%S}.png",
            full_page=True,
        )
    except Exception:
        pass


def processar_conta(
    browser: Browser,
    credencial: Credencial,
    vigencia_escolhida: str,
    pasta_downloads: Path,
    pasta_logs: Path,
    logger: logging.Logger,
    dados: DadosExecucao,
) -> ResultadoConta:
    resultado = ResultadoConta(
        linha_excel=credencial.linha_excel,
        usuario_mascarado=mascarar_usuario(credencial.usuario),
        contrato=credencial.contrato,
    )
    ultimo_erro = ""
    for tentativa in range(1, MAXIMO_TENTATIVAS_CONTA + 1):
        contexto: BrowserContext | None = None
        page: Page | None = None
        try:
            logger.info(
                "Linha %s | usuário %s | tentativa %s/%s.",
                credencial.linha_excel,
                resultado.usuario_mascarado,
                tentativa,
                MAXIMO_TENTATIVAS_CONTA,
            )
            contexto_conta = {
                "linha_excel": credencial.linha_excel,
                "usuario": resultado.usuario_mascarado,
                "tentativa": f"{tentativa}/{MAXIMO_TENTATIVAS_CONTA}",
            }
            with etapa_registrada(logger, "CRIAR_CONTEXTO_NAVEGADOR", **contexto_conta):
                contexto = browser.new_context(accept_downloads=True, no_viewport=True)
                contexto.set_default_timeout(TIMEOUT_PADRAO_MS)
                contexto.add_init_script(SCRIPT_CAPTURA_BLOB)
            with etapa_registrada(logger, "ABRIR_PORTAL_EMPRESA", **contexto_conta):
                page = abrir_portal_empresa(contexto, logger)
            with etapa_registrada(logger, "LOGIN_PORTAL", **contexto_conta):
                efetuar_login(page, credencial, logger)
            with etapa_registrada(logger, "NAVEGAR_SEGUNDA_VIA_BOLETOS", **contexto_conta):
                navegar_segunda_via(page, logger)
            with etapa_registrada(logger, "EXTRAIR_TABELA_BOLETOS", **contexto_conta):
                registros_extraidos = extrair_tabela(page, credencial, vigencia_escolhida)
                dados.registros_portal.extend(registros_extraidos)
                logger.info(
                    "ETAPA_ROBO | STATUS=DETALHE | ETAPA=EXTRAIR_TABELA_BOLETOS | REGISTROS=%s | LINHA_EXCEL=%s",
                    len(registros_extraidos),
                    credencial.linha_excel,
                )

            with etapa_registrada(logger, "IDENTIFICAR_LINHAS_PORTAL", **contexto_conta):
                linhas_portal = linhas_presentes_na_tabela(page)
            if not linhas_portal:
                if aviso_sem_registros(page):
                    resultado.status = "SEM_REGISTROS"
                    resultado.detalhe = "O portal informou que não há boletos para os filtros exibidos."
                    logger.info(
                        "Linha %s: nenhum boleto disponível para os filtros exibidos.",
                        credencial.linha_excel,
                    )
                    return resultado
                raise RuntimeError("A tabela foi exibida, mas nenhuma linha com Código e Número foi encontrada.")
            logger.info(
                "ETAPA_ROBO | STATUS=DETALHE | ETAPA=IDENTIFICAR_LINHAS_PORTAL | TOTAL=%s | LINHAS=%s | LINHA_EXCEL=%s",
                len(linhas_portal),
                ", ".join(f"{dados.codigo}/{dados.numero}" for dados in linhas_portal),
                credencial.linha_excel,
            )

            encontrou_algum = False
            for dados_linha in linhas_portal:
                codigo = dados_linha.codigo
                with etapa_registrada(
                    logger,
                    "LOCALIZAR_LINHA_TABELA",
                    **contexto_conta,
                    codigo=codigo,
                    numero=dados_linha.numero,
                ):
                    linha = localizar_linha_boleto(page, dados_linha)
                if linha is None:
                    logger.info(
                        "ETAPA_ROBO | STATUS=NAO_ENCONTRADO | ETAPA=LOCALIZAR_LINHA_TABELA | CODIGO=%s | NUMERO=%s | LINHA_EXCEL=%s",
                        codigo,
                        dados_linha.numero,
                        credencial.linha_excel,
                    )
                    continue
                encontrou_algum = True
                resultado.codigos_localizados += 1
                vigencia_codigo = credencial.vigencia or vigencia_escolhida
                pasta_codigo = pasta_downloads / nome_seguro(vigencia_codigo) / nome_seguro(codigo)
                pasta_codigo.mkdir(parents=True, exist_ok=True)
                logger.info(
                    "Código %s | Número %s localizado; iniciando downloads.",
                    codigo,
                    dados_linha.numero,
                )

                with etapa_registrada(
                    logger,
                    "PROCESSAR_NOTAS_CODIGO",
                    **contexto_conta,
                    codigo=codigo,
                ):
                    notas, ignoradas, falhas_nf = baixar_notas_codigo(
                        page,
                        contexto,
                        linha,
                        credencial,
                        codigo,
                        pasta_codigo,
                        logger,
                        dados,
                    )
                resultado.notas_baixadas += notas
                resultado.ignorados += ignoradas
                resultado.falhas += falhas_nf

                try:
                    with etapa_registrada(
                        logger,
                        "PROCESSAR_BOLETO_CODIGO",
                        **contexto_conta,
                        codigo=codigo,
                    ):
                        status_boleto, existente_boleto = baixar_boleto_codigo(
                            page,
                            linha,
                            codigo,
                            dados_linha.numero,
                            dados_linha.vencimento,
                            dados_linha.valor,
                            credencial.contrato,
                            pasta_codigo,
                            logger,
                        )
                    if status_boleto == "INDISPONIVEL":
                        resultado.indisponiveis += 1
                    elif existente_boleto:
                        resultado.ignorados += 1
                    else:
                        resultado.boletos_baixados += 1
                except Exception as exc:
                    resultado.falhas += 1
                    logger.exception(
                        "ETAPA_ROBO | STATUS=ERRO | ETAPA=PROCESSAR_BOLETO | CODIGO=%s | TIPO_ERRO=%s | ERRO=%s",
                        codigo,
                        type(exc).__name__,
                        exc,
                    )

                tipos = (
                    ("coparticipacao", "Coparticipacao", ".zip"),
                    ("coparticipacao_csv", "Coparticipacao_CSV", ".csv"),
                )
                for tipo_relatorio, descricao, extensao in tipos:
                    try:
                        status, existente = baixar_relatorio_icone(
                            page,
                            linha,
                            tipo_relatorio,
                            codigo,
                            descricao,
                            extensao,
                            credencial.contrato,
                            pasta_codigo,
                            logger,
                        )
                        if status == "INDISPONIVEL":
                            resultado.indisponiveis += 1
                        elif existente:
                            resultado.ignorados += 1
                        else:
                            resultado.relatorios_baixados += 1
                    except Exception as exc:
                        resultado.falhas += 1
                        logger.exception(
                            "ETAPA_ROBO | STATUS=ERRO | ETAPA=PROCESSAR_RELATORIO | CODIGO=%s | RELATORIO=%s | TIPO_ERRO=%s | ERRO=%s",
                            codigo,
                            descricao,
                            type(exc).__name__,
                            exc,
                        )

                with etapa_registrada(
                    logger,
                    "PROCESSAR_BOLETAS_CSV_E_PDF",
                    **contexto_conta,
                    codigo=codigo,
                ):
                    baixados_boletas, ignorados_boletas, indisponiveis_boletas, falhas_boletas = (
                        baixar_boletas_modal(
                            page,
                            linha,
                            codigo,
                            credencial.contrato,
                            pasta_codigo,
                            logger,
                        )
                    )
                resultado.relatorios_baixados += baixados_boletas
                resultado.ignorados += ignorados_boletas
                resultado.indisponiveis += indisponiveis_boletas
                resultado.falhas += falhas_boletas

            if not encontrou_algum:
                raise RuntimeError(
                    "Nenhum dos códigos da vigência selecionada foi encontrado na tabela desta conta."
                )
            resultado.status = "SUCESSO" if resultado.falhas == 0 else "PARCIAL"
            resultado.detalhe = "Processamento concluído."
            return resultado
        except Exception as exc:
            ultimo_erro = str(exc)
            logger.exception(
                "ETAPA_ROBO | STATUS=ERRO | ETAPA=PROCESSAR_CONTA | LINHA_EXCEL=%s | TENTATIVA=%s/%s | TIPO_ERRO=%s | ERRO=%s",
                credencial.linha_excel,
                tentativa,
                MAXIMO_TENTATIVAS_CONTA,
                type(exc).__name__,
                exc,
            )
            salvar_evidencia(page, pasta_logs, credencial.linha_excel)
            if tentativa < MAXIMO_TENTATIVAS_CONTA:
                time.sleep(2)
        finally:
            if contexto is not None:
                try:
                    contexto.close()
                except Exception:
                    pass
    resultado.status = "ERRO"
    resultado.falhas += 1
    resultado.detalhe = ultimo_erro
    return resultado


def ajustar_planilha(ws) -> None:
    preenchimento = PatternFill("solid", fgColor="1F4E78")
    fonte = Font(color="FFFFFF", bold=True)
    for celula in ws[1]:
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for coluna in ws.columns:
        largura = min(max(len(texto(celula.value)) for celula in coluna) + 2, 45)
        ws.column_dimensions[coluna[0].column_letter].width = max(largura, 12)


def adicionar_aba_registros(wb: Workbook, nome: str, registros: list[dict[str, str]]) -> None:
    ws = wb.create_sheet(nome)
    if not registros:
        ws.append(["INFORMACAO"])
        ws.append(["Nenhum registro coletado."])
        ajustar_planilha(ws)
        return
    colunas: list[str] = []
    for registro in registros:
        for chave in registro:
            if chave not in colunas:
                colunas.append(chave)
    ws.append(colunas)
    for registro in registros:
        ws.append([registro.get(coluna, "") for coluna in colunas])
    ajustar_planilha(ws)


def salvar_planilha_resultados(
    dados: DadosExecucao,
    resultados: list[ResultadoConta],
    pasta_downloads: Path,
    vigencia: str,
) -> Path:
    caminho = pasta_downloads / f"Informacoes_Segunda_Via_Boletos_{nome_seguro(vigencia)}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb = Workbook()
    ws_resumo = wb.active
    ws_resumo.title = "RESUMO"
    cabecalhos = [
        "LINHA_EXCEL",
        "USUARIO_MASCARADO",
        "CONTRATO",
        "STATUS",
        "CODIGOS_LOCALIZADOS",
        "NOTAS_BAIXADAS",
        "BOLETOS_BAIXADOS",
        "RELATORIOS_BAIXADOS",
        "IGNORADOS",
        "INDISPONIVEIS",
        "FALHAS",
        "DETALHE",
    ]
    ws_resumo.append(cabecalhos)
    for r in resultados:
        ws_resumo.append(
            [
                r.linha_excel,
                r.usuario_mascarado,
                r.contrato,
                r.status,
                r.codigos_localizados,
                r.notas_baixadas,
                r.boletos_baixados,
                r.relatorios_baixados,
                r.ignorados,
                r.indisponiveis,
                r.falhas,
                r.detalhe,
            ]
        )
    ajustar_planilha(ws_resumo)
    adicionar_aba_registros(wb, "SEGUNDA_VIA_BOLETOS", dados.registros_portal)
    adicionar_aba_registros(wb, "NOTAS_FISCAIS", dados.registros_notas)
    wb.save(caminho)
    return caminho


def salvar_resumo_csv(resultados: list[ResultadoConta], pasta_logs: Path) -> Path:
    caminho = pasta_logs / f"resumo_{datetime.now():%Y%m%d_%H%M%S}.csv"
    with caminho.open("w", newline="", encoding="utf-8-sig") as arquivo:
        writer = csv.writer(arquivo, delimiter=";")
        writer.writerow(
            [
                "LINHA_EXCEL",
                "USUARIO_MASCARADO",
                "CONTRATO",
                "STATUS",
                "CODIGOS_LOCALIZADOS",
                "NOTAS_BAIXADAS",
                "BOLETOS_BAIXADOS",
                "RELATORIOS_BAIXADOS",
                "IGNORADOS",
                "INDISPONIVEIS",
                "FALHAS",
                "DETALHE",
            ]
        )
        for r in resultados:
            writer.writerow(
                [
                    r.linha_excel,
                    r.usuario_mascarado,
                    r.contrato,
                    r.status,
                    r.codigos_localizados,
                    r.notas_baixadas,
                    r.boletos_baixados,
                    r.relatorios_baixados,
                    r.ignorados,
                    r.indisponiveis,
                    r.falhas,
                    r.detalhe,
                ]
            )
    return caminho


def executar_robo(
    pasta_destino: Path,
    planilha_selecionada: Path | None,
    vigencia: str,
    fila_log: queue.Queue[str],
    callback_final,
) -> None:
    base = pasta_do_script()
    pasta_logs = base / NOME_PASTA_LOGS
    logger, arquivo_log, arquivo_erros = configurar_logger(pasta_logs, fila_log)
    resultados: list[ResultadoConta] = []
    dados = DadosExecucao()
    inicio = time.monotonic()
    try:
        logger.info(
            "ETAPA_ROBO | STATUS=INICIO | ETAPA=EXECUCAO_COMPLETA | VIGENCIA=%s | DESTINO=%s",
            vigencia,
            pasta_destino,
        )
        with etapa_registrada(logger, "LOCALIZAR_E_VALIDAR_PLANILHA", vigencia=vigencia):
            if planilha_selecionada is not None:
                planilha = planilha_selecionada.expanduser().resolve()
                if not planilha.is_file():
                    raise FileNotFoundError(f"A planilha selecionada não existe: {planilha}")
                if planilha.suffix.lower() not in {".xlsx", ".xlsm"}:
                    raise ValueError(
                        f"Formato não suportado: {planilha.suffix or 'sem extensão'}. "
                        "Selecione uma planilha .xlsx ou .xlsm."
                    )
            else:
                planilha = localizar_planilha(base)
        with etapa_registrada(
            logger,
            "LER_CREDENCIAIS_PLANILHA",
            vigencia=vigencia,
            planilha=planilha,
        ):
            credenciais = ler_credenciais(planilha, vigencia)
        logger.info("Planilha de controle: %s", planilha)
        logger.info("Vigência selecionada: %s", vigencia)
        logger.info("Linhas ativas a processar: %s", len(credenciais))

        with sync_playwright() as playwright:
            opcoes_navegador = {
                "headless": not EXECUTAR_VISIVEL,
                "args": ["--start-maximized"],
            }
            with etapa_registrada(logger, "ABRIR_NAVEGADOR"):
                if PREFERIR_MICROSOFT_EDGE:
                    try:
                        browser = playwright.chromium.launch(channel="msedge", **opcoes_navegador)
                        logger.info("Navegador utilizado: Microsoft Edge.")
                    except Exception as exc:
                        logger.warning("Microsoft Edge indisponível; usando Chromium: %s", exc)
                        browser = playwright.chromium.launch(**opcoes_navegador)
                else:
                    browser = playwright.chromium.launch(**opcoes_navegador)
            try:
                for posicao, credencial in enumerate(credenciais, start=1):
                    logger.info(
                        "Processando %s/%s | linha %s | faltam %s.",
                        posicao,
                        len(credenciais),
                        credencial.linha_excel,
                        len(credenciais) - posicao,
                    )
                    with etapa_registrada(
                        logger,
                        "PROCESSAR_LINHA_PLANILHA",
                        linha_excel=credencial.linha_excel,
                        posicao=f"{posicao}/{len(credenciais)}",
                        usuario=mascarar_usuario(credencial.usuario),
                    ):
                        resultados.append(
                            processar_conta(
                                browser,
                                credencial,
                                vigencia,
                                pasta_destino,
                                pasta_logs,
                                logger,
                                dados,
                            )
                        )
            finally:
                with etapa_registrada(logger, "FECHAR_NAVEGADOR"):
                    browser.close()

        with etapa_registrada(logger, "SALVAR_PLANILHA_RESULTADOS", vigencia=vigencia):
            planilha_saida = salvar_planilha_resultados(dados, resultados, pasta_destino, vigencia)
        with etapa_registrada(logger, "SALVAR_RESUMO_CSV", vigencia=vigencia):
            resumo_csv = salvar_resumo_csv(resultados, pasta_logs)
        duracao = int(time.monotonic() - inicio)
        logger.info(
            "FIM | sucesso=%s | parcial=%s | erro=%s | NFs=%s | boletos=%s | relatórios=%s | ignorados=%s | indisponíveis=%s | falhas=%s | tempo=%02d:%02d:%02d",
            sum(r.status == "SUCESSO" for r in resultados),
            sum(r.status == "PARCIAL" for r in resultados),
            sum(r.status == "ERRO" for r in resultados),
            sum(r.notas_baixadas for r in resultados),
            sum(r.boletos_baixados for r in resultados),
            sum(r.relatorios_baixados for r in resultados),
            sum(r.ignorados for r in resultados),
            sum(r.indisponiveis for r in resultados),
            sum(r.falhas for r in resultados),
            duracao // 3600,
            (duracao % 3600) // 60,
            duracao % 60,
        )
        logger.info(
            "ETAPA_ROBO | STATUS=SUCESSO | ETAPA=EXECUCAO_COMPLETA | DURACAO_SEG=%s | VIGENCIA=%s",
            duracao,
            vigencia,
        )
        callback_final(
            True,
            f"Execução concluída.\n\nPlanilha: {planilha_saida}\nResumo CSV: {resumo_csv}"
            f"\nLog geral: {arquivo_log}\nLog de erros e diagnóstico PDF: {arquivo_erros}",
        )
    except Exception as exc:
        logger.exception(
            "ETAPA_ROBO | STATUS=ERRO | ETAPA=EXECUCAO_COMPLETA | TIPO_ERRO=%s | ERRO=%s",
            type(exc).__name__,
            exc,
        )
        callback_final(
            False,
            f"Execução interrompida:\n\n{exc}\n\nLog geral: {arquivo_log}"
            f"\nLog de erros e diagnóstico PDF: {arquivo_erros}",
        )


class Aplicacao:
    def __init__(self) -> None:
        self.root = Tk()
        self.root.title("Robô Nordeste Saúde - Notas, Relatórios e Boletos")
        self.root.geometry("940x720")
        self.root.minsize(780, 580)
        self.fila_log: queue.Queue[str] = queue.Queue()
        planilha_automatica = procurar_planilha_existente(pasta_do_script())
        self.planilha_var = StringVar(value=str(planilha_automatica) if planilha_automatica else "")
        self.pasta_var = StringVar(value=str(pasta_do_script() / NOME_PASTA_DOWNLOADS))
        self.vigencia_var = StringVar(value="VIG 20")

        topo = Frame(self.root, padx=12, pady=12)
        topo.pack(fill=X)

        Label(topo, text="Planilha de controle:").pack(anchor="w")
        linha_planilha = Frame(topo)
        linha_planilha.pack(fill=X, pady=(5, 10))
        self.entrada_planilha = Entry(linha_planilha, textvariable=self.planilha_var)
        self.entrada_planilha.pack(side=LEFT, fill=X, expand=True)
        Button(
            linha_planilha,
            text="Selecionar planilha...",
            command=self.selecionar_planilha,
        ).pack(side=RIGHT, padx=(8, 0))

        Label(topo, text="Pasta para salvar os downloads:").pack(anchor="w")
        linha_pasta = Frame(topo)
        linha_pasta.pack(fill=X, pady=(5, 10))
        self.entrada_pasta = Entry(linha_pasta, textvariable=self.pasta_var)
        self.entrada_pasta.pack(side=LEFT, fill=X, expand=True)
        Button(linha_pasta, text="Selecionar...", command=self.selecionar_pasta).pack(side=RIGHT, padx=(8, 0))

        linha_vigencia = Frame(topo)
        linha_vigencia.pack(fill=X)
        Label(linha_vigencia, text="Vigência para download:").pack(side=LEFT)
        self.combo_vigencia = Combobox(
            linha_vigencia,
            textvariable=self.vigencia_var,
            values=("VIG 01", "VIG 10", "VIG 20", "TODAS"),
            state="readonly",
            width=14,
        )
        self.combo_vigencia.pack(side=LEFT, padx=(8, 0))

        botoes = Frame(self.root, padx=12)
        botoes.pack(fill=X)
        self.botao_iniciar = Button(botoes, text="Iniciar downloads", command=self.iniciar, height=2)
        self.botao_iniciar.pack(side=LEFT)
        Button(botoes, text="Abrir pasta do script", command=self.abrir_pasta_script, height=2).pack(side=LEFT, padx=8)

        self.status_var = StringVar(value="Aguardando início.")
        Label(self.root, textvariable=self.status_var, anchor="w", padx=12, pady=8).pack(fill=X)
        self.console = ScrolledText(self.root, state="disabled", wrap="word", font=("Consolas", 9))
        self.console.pack(fill=BOTH, expand=True, padx=12, pady=(0, 12))
        self.root.after(150, self.atualizar_console)

    def selecionar_planilha(self) -> None:
        caminho_atual = self.planilha_var.get().strip()
        pasta_inicial = Path(caminho_atual).parent if caminho_atual else pasta_do_script()
        selecionada = filedialog.askopenfilename(
            title="Selecione a planilha de controle",
            initialdir=str(pasta_inicial),
            filetypes=[
                ("Planilhas do Excel", "*.xlsx *.xlsm"),
                ("Excel XLSX", "*.xlsx"),
                ("Excel XLSM", "*.xlsm"),
                ("Todos os arquivos", "*.*"),
            ],
        )
        if selecionada:
            self.planilha_var.set(selecionada)

    def selecionar_pasta(self) -> None:
        selecionada = filedialog.askdirectory(initialdir=self.pasta_var.get() or str(pasta_do_script()))
        if selecionada:
            self.pasta_var.set(selecionada)

    def abrir_pasta_script(self) -> None:
        caminho = str(pasta_do_script())
        try:
            if os.name == "nt":
                os.startfile(caminho)  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                os.system(f"open {caminho!r}")
            else:
                os.system(f"xdg-open {caminho!r}")
        except Exception as exc:
            messagebox.showerror("Erro", str(exc))

    def iniciar(self) -> None:
        planilha_texto = self.planilha_var.get().strip()
        destino_texto = self.pasta_var.get().strip()
        vigencia = self.vigencia_var.get().strip()
        if not planilha_texto:
            planilha_automatica = procurar_planilha_existente(pasta_do_script())
            if planilha_automatica is None:
                messagebox.showwarning("Atenção", "Selecione a planilha de controle.")
                return
            planilha = planilha_automatica
            self.planilha_var.set(str(planilha))
        else:
            planilha = Path(planilha_texto).expanduser()
        if not planilha.is_file():
            messagebox.showerror("Erro", f"A planilha selecionada não existe:\n{planilha}")
            return
        if planilha.suffix.lower() not in {".xlsx", ".xlsm"}:
            messagebox.showerror("Erro", "Selecione uma planilha do Excel no formato .xlsx ou .xlsm.")
            return
        if not destino_texto or vigencia not in {"VIG 01", "VIG 10", "VIG 20", "TODAS"}:
            messagebox.showwarning("Atenção", "Selecione a pasta de destino e uma vigência válida.")
            return
        destino = Path(destino_texto).expanduser()
        try:
            destino.mkdir(parents=True, exist_ok=True)
        except Exception as exc:
            messagebox.showerror("Erro", f"Não foi possível criar a pasta:\n{exc}")
            return
        self.botao_iniciar.config(state="disabled")
        self.combo_vigencia.config(state="disabled")
        self.entrada_planilha.config(state="disabled")
        self.status_var.set(f"Executando {vigencia}. Não feche o navegador.")
        threading.Thread(
            target=executar_robo,
            args=(destino, planilha, vigencia, self.fila_log, self.finalizar_threadsafe),
            daemon=True,
        ).start()

    def finalizar_threadsafe(self, sucesso: bool, mensagem: str) -> None:
        self.root.after(0, lambda: self.finalizar(sucesso, mensagem))

    def finalizar(self, sucesso: bool, mensagem: str) -> None:
        self.botao_iniciar.config(state="normal")
        self.combo_vigencia.config(state="readonly")
        self.entrada_planilha.config(state="normal")
        self.status_var.set("Concluído." if sucesso else "Execução interrompida.")
        (messagebox.showinfo if sucesso else messagebox.showerror)("Concluído" if sucesso else "Erro", mensagem)

    def atualizar_console(self) -> None:
        try:
            while True:
                mensagem = self.fila_log.get_nowait()
                self.console.config(state="normal")
                self.console.insert(END, mensagem + "\n")
                self.console.see(END)
                self.console.config(state="disabled")
        except queue.Empty:
            pass
        self.root.after(150, self.atualizar_console)

    def executar(self) -> None:
        self.root.mainloop()


def main() -> None:
    Aplicacao().executar()


if __name__ == "__main__":
    main()
